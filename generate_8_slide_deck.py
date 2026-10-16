import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette: Premium Deep Navy, Vibrant Teal Accent, Slate Gray, Pure White
    BG_COLOR = RGBColor(11, 19, 43)        # #0b132b deep navy
    CARD_BG = RGBColor(28, 37, 65)         # #1c2541 card background
    TEXT_LIGHT = RGBColor(240, 244, 248)   # clean off-white
    TEXT_MUTED = RGBColor(160, 174, 192)   # slate gray
    ACCENT_CYAN = RGBColor(79, 209, 197)   # teal/cyan accent
    ACCENT_BLUE = RGBColor(66, 153, 225)   # bright blue
    ACCENT_GREEN = RGBColor(72, 187, 120)  # success green

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.color.rgb = BG_COLOR
        return bg

    def add_header(slide, title_text, category="AEROSTORE PLATFORM PRESENTATION"):
        # Category pill
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_CYAN

        # Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    t_box1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(3.5))
    tf1 = t_box1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "CAPGEMINI L1 ENGINEERING USE CASES"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_CYAN

    p1 = tf1.add_paragraph()
    p1.text = "AeroStore Enterprise E-Commerce"
    p1.font.size = Pt(38)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT

    p2 = tf1.add_paragraph()
    p2.text = "Full-Stack N-Tier Platform Powered by Grounded RAG & Enterprise Reliability"
    p2.font.size = Pt(20)
    p2.font.color.rgb = ACCENT_BLUE

    p3 = tf1.add_paragraph()
    p3.text = "\n• React 18 UI  |  FastAPI Layered Backend  |  In-Memory RAG  |  SQL Server (SSMS)\n• Built for Mentor Evaluation & Enterprise Architecture Defense"
    p3.font.size = Pt(14)
    p3.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 2: Project Overview & Core Problem
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Project Overview & Problem Statement", "01. INTRODUCTION")

    cards2 = [
        ("The Industry Problem", 
         "• Standard e-commerce demos are toy CRUD apps that fail in production.\n• Generic LLM chatbots hallucinate return terms and inventory data.\n• Concurrent users trigger race conditions and double charges during flash sales.",
         ACCENT_CYAN),
        ("The AeroStore Solution", 
         "• Production-grade 5-Tier decoupled architecture.\n• Retrieval-Augmented Generation (RAG) for verifiable, hallucination-free store guidance.\n• Built-in concurrency tokens and idempotent payment sagas.",
         ACCENT_BLUE),
        ("Target Outcomes", 
         "• Sub-10ms localized vector retrieval with zero heavy C-dependencies.\n• 100% reliable cart synchronization with HTTP 409 conflict safety.\n• Complete, audit-ready AI engineering lifecycle (Use Cases 1–5).",
         ACCENT_GREEN)
    ]

    for i, (title, content, color) in enumerate(cards2):
        x = Inches(0.8 + i * 3.95)
        shape = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(3.7), Inches(4.7))
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = color
        shape.line.width = Pt(1.5)

        tb = s2.shapes.add_textbox(x + Inches(0.25), Inches(2.2), Inches(3.2), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(18)
        pt.font.bold = True
        pt.font.color.rgb = color

        pc = tf.add_paragraph()
        pc.text = "\n" + content
        pc.font.size = Pt(13)
        pc.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 3: What is RAG & Why Did We Implement It?
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "RAG (Retrieval-Augmented Generation): The Core AI Engine", "02. AI CAPABILITIES")

    cards3 = [
        ("What is RAG?",
         "• Combines semantic document search with LLM synthesis.\n• Ingests store policies, product hardware specs, and ADR decisions.\n• Vector cosine similarity fetches the top-k relevant knowledge chunks before the AI answers.",
         ACCENT_CYAN),
        ("Why Implement RAG?",
         "• Prevents Hallucinations: Generic AI invents refund rules; RAG constrains answers strictly to store facts.\n• Zero Retraining Cost: When policies change, you update text—not expensive model weights.\n• Verifiable Citations: Answers cite source titles and relevance scores.",
         ACCENT_BLUE),
        ("Live RAG In Action",
         "• `/rag return policy` &rarr; 30-day RMA rules, packaging, restocking terms.\n• `/rag warranty` &rarr; 2-year manufacturer hardware coverage.\n• `/rag headphone specs` &rarr; 40mm titanium drivers, 42dB hybrid ANC.\n• `/rag why microservices` &rarr; Ingests ADR-001 architectural decisions.",
         ACCENT_GREEN)
    ]

    for i, (title, content, color) in enumerate(cards3):
        x = Inches(0.8 + i * 3.95)
        shape = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(3.7), Inches(4.7))
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = color
        shape.line.width = Pt(1.5)

        tb = s3.shapes.add_textbox(x + Inches(0.25), Inches(2.2), Inches(3.2), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(18)
        pt.font.bold = True
        pt.font.color.rgb = color

        pc = tf.add_paragraph()
        pc.text = "\n" + content
        pc.font.size = Pt(13)
        pc.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 4: What Makes AeroStore Different from Others?
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Competitive Edge: What Makes AeroStore Different?", "03. KEY DIFFERENTIATORS")

    diff_cards = [
        ("1. Strict Grounded RAG", "Not a generic chatbot wrapper. Answers are backed by cosine similarity and verifiable knowledge citations.", ACCENT_CYAN),
        ("2. Flash-Sale OCC Concurrency", "Optimistic Concurrency Control with integer tokens rejects stale cart overwrites via HTTP 409 Conflict.", ACCENT_BLUE),
        ("3. Idempotent Payment Sagas", "Guarantees zero double-charging even if users double-click checkout or network drops midway.", ACCENT_GREEN),
        ("4. True N-Tier Enterprise Code", "Rigid separation across 5 layers (UI &rarr; API &rarr; BLL &rarr; DAL &rarr; DB). Zero monolithic coupling.", ACCENT_CYAN),
        ("5. Dual Storage Strategy", "Zero-setup in-memory dev database + production-ready Microsoft SQL Server (SSMS) relational schemas.", ACCENT_BLUE),
        ("6. 5 Full AI Engineering Cycles", "Automates architecture review, contract-first OpenAPI, WCAG accessibility, static code audit, and stress testing.", ACCENT_GREEN)
    ]

    for idx, (title, desc, color) in enumerate(diff_cards):
        row = idx // 3
        col = idx % 3
        x = Inches(0.8 + col * 3.95)
        y = Inches(2.0 + row * 2.45)

        shape = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.7), Inches(2.2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = color
        shape.line.width = Pt(1.2)

        tb = s4.shapes.add_textbox(x + Inches(0.2), y + Inches(0.15), Inches(3.3), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(15)
        pt.font.bold = True
        pt.font.color.rgb = color

        pc = tf.add_paragraph()
        pc.text = desc
        pc.font.size = Pt(12)
        pc.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 5: 5-Tier System Architecture
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Enterprise N-Tier System Architecture", "04. SYSTEM DESIGN")

    tiers = [
        ("Tier 1: Presentation (React 18 + Vite)", "Modern glassmorphic UI, live Cart Drawer, AeroBot AI Concierge widget, role-based dashboards.", ACCENT_CYAN),
        ("Tier 2: API & Gateway (FastAPI)", "REST controllers, JWT security, Swagger/OpenAPI docs, and RFC 7807 Problem Details error middleware.", ACCENT_BLUE),
        ("Tier 3: Business Logic Tier (BLL)", "Pure business rules: ProductService, CartService, PaymentService, CouponService, and RagEngine.", ACCENT_GREEN),
        ("Tier 4: Data Access Layer (DAL)", "Repository pattern abstracting SQL and in-memory queries (ProductRepository, OrderRepository).", ACCENT_CYAN),
        ("Tier 5: Database Persistence", "In-memory mock store for rapid dev + Enterprise Microsoft SQL Server (SSMS) with indexed relational tables.", ACCENT_BLUE)
    ]

    for i, (t_name, t_desc, col) in enumerate(tiers):
        y = Inches(1.9 + i * 1.0)
        shape = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(0.85))
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = col
        shape.line.width = Pt(1.2)

        tb = s5.shapes.add_textbox(Inches(1.0), y + Inches(0.1), Inches(11.3), Inches(0.65))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = t_name
        pt.font.size = Pt(15)
        pt.font.bold = True
        pt.font.color.rgb = col

        pc = tf.add_paragraph()
        pc.text = t_desc
        pc.font.size = Pt(12)
        pc.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 6: 5 AI Engineering Use Cases (L1 Scope)
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Complete 5 L1 AI Engineering Use Cases", "05. AI LIFECYCLE")

    use_cases = [
        ("UC1: Architecture Synthesis", "Evaluates Monolith vs Microservices tradeoffs, DB-per-service isolation, and creates ADR-001/002/003."),
        ("UC2: Contract-First Scaffolding", "Generates OpenAPI 3.0 contracts and scaffolds FastAPI controllers with RFC 7807 error models."),
        ("UC3: Frontend Risk & A11y", "Automates WCAG 2.1 AA audits, checking ARIA labels, color contrast, and keyboard navigation."),
        ("UC4: Code & Security Review", "Static AST review detecting SQL injection risks, exposed credentials, and missing OCC version checks."),
        ("UC5: Edge Case Testing & Benchmark", "Concurrent multi-threaded stress tests for cart collisions, idempotency keys, and retry reprompting.")
    ]

    for i, (uc_title, uc_desc) in enumerate(use_cases):
        y = Inches(1.9 + i * 1.0)
        shape = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(0.85))
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = ACCENT_CYAN if i % 2 == 0 else ACCENT_BLUE
        shape.line.width = Pt(1.2)

        tb = s6.shapes.add_textbox(Inches(1.0), y + Inches(0.1), Inches(11.3), Inches(0.65))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = uc_title
        pt.font.size = Pt(15)
        pt.font.bold = True
        pt.font.color.rgb = ACCENT_CYAN if i % 2 == 0 else ACCENT_BLUE

        pc = tf.add_paragraph()
        pc.text = uc_desc
        pc.font.size = Pt(12)
        pc.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 7: Business Benefits & Value Delivered
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Business Value & ROI for Enterprise Stakeholders", "06. IMPACT")

    b_cards = [
        ("Customer Experience",
         "• Sub-second answers to warranty & return policies.\n• Zero checkout errors or duplicate card debits.\n• Intuitive dark/light UI with instant cart drawer updates.",
         ACCENT_CYAN),
        ("Business & Operations",
         "• Up to 60% reduction in customer support tickets via RAG self-serve answers.\n• Protected flash-sale revenue with zero overselling.\n• Verifiable audit trail with citation tracking on AI advice.",
         ACCENT_BLUE),
        ("Engineering & Scalability",
         "• Plug-and-play N-Tier modules make onboarding fast.\n• Clean test harness prevents regressions.\n• Smooth transition path from prototype to enterprise SQL Server.",
         ACCENT_GREEN)
    ]

    for i, (title, content, color) in enumerate(b_cards):
        x = Inches(0.8 + i * 3.95)
        shape = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(3.7), Inches(4.7))
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = color
        shape.line.width = Pt(1.5)

        tb = s7.shapes.add_textbox(x + Inches(0.25), Inches(2.2), Inches(3.2), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(18)
        pt.font.bold = True
        pt.font.color.rgb = color

        pc = tf.add_paragraph()
        pc.text = "\n" + content
        pc.font.size = Pt(13)
        pc.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 8: Live Demonstration & Summary
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Live Demonstration Flow & Conclusion", "07. SUMMARY & DEMO")

    cards8 = [
        ("Demo Step 1: E-Commerce Storefront",
         "• Browse live catalog across 6 categories.\n• Add items to cart & inspect responsive state.",
         ACCENT_CYAN),
        ("Demo Step 2: Live RAG AI Concierge",
         "• Query `/rag return policy` & `/rag headphone specs`.\n• Inspect verifiable source citations & cosine scores.",
         ACCENT_BLUE),
        ("Demo Step 3: Concurrency & Backend Docs",
         "• Show HTTP 409 OCC cart conflict handling.\n• Review interactive FastAPI Swagger at `/docs`.",
         ACCENT_GREEN)
    ]

    for i, (title, content, color) in enumerate(cards8):
        x = Inches(0.8 + i * 3.95)
        shape = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(3.7), Inches(3.2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = color
        shape.line.width = Pt(1.5)

        tb = s8.shapes.add_textbox(x + Inches(0.25), Inches(2.2), Inches(3.2), Inches(2.8))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(16)
        pt.font.bold = True
        pt.font.color.rgb = color

        pc = tf.add_paragraph()
        pc.text = "\n" + content
        pc.font.size = Pt(13)
        pc.font.color.rgb = TEXT_LIGHT

    # Bottom summary box
    sum_shape = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.5), Inches(11.7), Inches(1.3))
    sum_shape.fill.solid()
    sum_shape.fill.fore_color.rgb = CARD_BG
    sum_shape.line.color.rgb = ACCENT_CYAN
    sum_shape.line.width = Pt(1.2)

    sum_tb = s8.shapes.add_textbox(Inches(1.0), Inches(5.6), Inches(11.3), Inches(1.1))
    sum_tf = sum_tb.text_frame
    sum_tf.word_wrap = True
    p_sum = sum_tf.paragraphs[0]
    p_sum.text = "🎯 Key Conclusion"
    p_sum.font.size = Pt(15)
    p_sum.font.bold = True
    p_sum.font.color.rgb = ACCENT_CYAN

    p_sum2 = sum_tf.add_paragraph()
    p_sum2.text = "AeroStore proves that modern AI is most effective when paired with resilient software engineering principles: strict N-Tier separation, optimistic concurrency, and grounded, verifiable RAG."
    p_sum2.font.size = Pt(13)
    p_sum2.font.color.rgb = TEXT_LIGHT

    # Save presentation
    output_path = os.path.join(os.path.dirname(__file__), "AeroStore_8_Slide_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_deck()
